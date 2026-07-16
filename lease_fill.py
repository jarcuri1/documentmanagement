"""
LeaseFill — fills a lease and queues it for approval
====================================================
Piece 3 of the pipeline (see HANDOFF_LEASE_AGENT.md). Takes a structured
intake, fills the landlord's OWN Word lease (single-family OR multifamily),
converts it to PDF with LibreOffice, and drops the completed (unsigned)
lease + a job file into Dropbox\\Leases\\Pending, then writes an approval
card onto the existing approvals rails so the supervisor can push it to your
phone.

Why fill the .docx: values replace the blanks and the text reflows
naturally -- no coordinate math, and adding/editing a template is trivial.
The legal language stays the landlord's own. Signature / initial / Print
Name / Date blocks are left untouched -- those are placed by Authentisign.

FAIL-CLOSED, twice over:
  * Every blank is found by the text around it. If a REQUIRED field's anchor
    is missing (the template changed), the whole job aborts -- a half-filled
    lease is never produced or queued.
  * LibreOffice exits 0 even when it fails to convert, so we verify the PDF
    actually appeared and is non-empty; if not, we abort.

SINGLE vs MULTI FAMILY: `lease_type` selects the template. Single-family has
utility checkboxes and septic/oil clauses; multifamily does not and instead
has an optional pets list. Both carry two tenant slots; `tenants` may hold
one or two people, each of whom becomes an Authentisign signer.

SENSITIVE DATA: SSNs, if supplied, are written ONLY onto the lease. They are
never written into the job file or the approval card (both of which leave
the machine). SSN is optional.

RUN:
  python lease_fill.py --intake intake.json          # fill + queue
  python lease_fill.py --intake intake.json --dry-run  # build PDF beside the
                        # intake, queue NOTHING

INTAKE FORMAT (intake.json):
{
  "lease_type": "single_family",          # or "multi_family"
  "landlord": "Premio Property Management LLC",
  "property": "123 Main St Apt 2, Waterbury CT",
  "premises_address": "123 Main St Apt 2, Waterbury, CT 06702",
  "agreement_date": "2026-07-16",         # ISO or free text; omit = today
  "term_start": "2026-08-01",
  "term_end":   "2027-07-31",
  "rent": "2,150",
  "deposit": "2,150",
  "utilities": { "water": "City", "wastewater": "Sewer", "fuel": "Oil" },  # single_family only
  "pets": ["Rex / Labrador / Black"],     # multi_family only, optional
  "tenants": [
    { "name": "John Smith", "email": "jsmith@example.com",
      "address": "123 Main St Apt 2", "city_state_zip": "Waterbury, CT 06702",
      "ssn": "XXX-XX-1234" }
  ]
}
"""

import argparse
import json
import os
import re
import shutil
import subprocess
import sys
import time
from datetime import date, datetime, timezone
from pathlib import Path

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

# ----------------------------------------------------------------------
# CONFIG — env-overridable; defaults mirror lease_watcher / the handoff
# ----------------------------------------------------------------------
_SHARED_ROOT = os.environ.get("LEASE_SHARED_ROOT", r"C:\AIAgents\shared")
_LEASES_ROOT = os.environ.get("LEASE_DROPBOX_ROOT", r"D:\Dropbox\Dropbox\Leases")
_TEMPLATE_DIR = Path(os.environ.get("LEASE_TEMPLATE_DIR", str(Path(__file__).with_name("templates"))))


def _p(env_key, *default_parts, root):
    v = os.environ.get(env_key)
    return Path(v) if v else Path(root, *default_parts)


def _find_soffice():
    """LibreOffice isn't on PATH after a default Windows install, so look in
    the standard install locations before falling back to the PATH name."""
    override = os.environ.get("LEASE_SOFFICE")
    if override:
        return override
    for c in (r"C:\Program Files\LibreOffice\program\soffice.exe",
              r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"):
        if Path(c).exists():
            return c
    return "soffice"


CONFIG = {
    "templates": {
        "single_family": _TEMPLATE_DIR / "single_family_lease.docx",
        "multi_family": _TEMPLATE_DIR / "multi_family_lease.docx",
    },
    "pending_dir":       _p("LEASE_PENDING_DIR", "Pending", root=_LEASES_ROOT),
    "intake_dir":        _p("LEASE_INTAKE_DIR", "Intake", root=_LEASES_ROOT),
    "pending_cards_dir": _p("LEASE_CARDS_DIR", "approvals", "pending", root=_SHARED_ROOT),
    "push_outbox_dir":   _p("LEASE_PUSH_OUTBOX", "push_outbox", root=_SHARED_ROOT),
    "dropbox_rel_root":  os.environ.get("LEASE_DROPBOX_REL_ROOT", "/Leases"),
    "dropbox_token":     os.environ.get("LEASE_DROPBOX_TOKEN", ""),
    "soffice_bin":       _find_soffice(),
    "convert_timeout_s": int(os.environ.get("LEASE_CONVERT_TIMEOUT", "120")),
}


class LeaseFillError(Exception):
    pass


_push_seq = 0


def push(title, body, data=None):
    """Drop a status push onto the Supervisor's push_outbox rail (swept every
    15s). Used only for fill FAILURES — a successful fill's approval card
    auto-pushes when it lands in approvals\\pending."""
    global _push_seq
    _push_seq += 1
    outbox = CONFIG["push_outbox_dir"]
    outbox.mkdir(parents=True, exist_ok=True)
    payload = {"title": title, "body": body, "data": {**(data or {}), "kind": "lease"}}
    name = f"lease-{os.getpid()}-{int(time.time() * 1000)}-{_push_seq}.json"
    tmp = outbox / (name + ".tmp")
    tmp.write_text(json.dumps(payload), encoding="utf-8")
    tmp.replace(outbox / name)


def make_pdf_url(pdf_path):
    """Create a Dropbox share URL for the filled lease so the approval card can
    offer a tap-to-open button on Jay's phone. Needs a Dropbox token
    (LEASE_DROPBOX_TOKEN) and the `dropbox` package. Degrades gracefully: if
    either is missing the card is still queued, just without the link."""
    token = CONFIG["dropbox_token"]
    if not token:
        return ""
    rel = f"{CONFIG['dropbox_rel_root'].rstrip('/')}/Pending/{Path(pdf_path).name}"
    try:
        import dropbox
        from dropbox.exceptions import ApiError
    except ImportError:
        print("WARN: `dropbox` package not installed — card will omit pdf_url "
              "(pip install dropbox)")
        return ""
    try:
        dbx = dropbox.Dropbox(token)
        try:
            return dbx.sharing_create_shared_link_with_settings(rel).url
        except ApiError:
            links = dbx.sharing_list_shared_links(path=rel, direct_only=True).links
            return links[0].url if links else ""
    except Exception as e:
        print(f"WARN: Dropbox share-link generation failed ({e}) — card omits pdf_url")
        return ""


# ----------------------------------------------------------------------
# DOCX fill primitives — operate on <w:t> text only, so <w:br/> line breaks
# and every other structural element are preserved untouched.
# ----------------------------------------------------------------------
_UND = re.compile(r"_{3,}")


def _text_nodes(paragraph):
    return list(paragraph._p.iter(qn("w:t")))


def fill_blanks(paragraph, values):
    """Replace the underscore-runs in `paragraph`, in order, with `values`.
    A None/"" value leaves that blank as-is. Returns how many blanks existed."""
    vi = 0
    for t in _text_nodes(paragraph):
        s = t.text or ""
        if "_" not in s:
            continue
        out, last = [], 0
        for m in _UND.finditer(s):
            out.append(s[last:m.start()])
            v = values[vi] if vi < len(values) else None
            vi += 1
            out.append(m.group(0) if v in (None, "") else str(v))
            last = m.end()
        out.append(s[last:])
        t.text = "".join(out)
    return vi


def insert_before_break(paragraph, text):
    """Put `text` on the line before the first <w:br/> (single-family tenant
    name sits on the numbered-list line, which has no underscore blank)."""
    for r in paragraph._p.iter(qn("w:r")):
        br = r.find(qn("w:br"))
        if br is not None:
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = text
            br.addprevious(t)
            return True
    raise LeaseFillError("expected a line break in the tenant block but found none")


def check_nth_box(paragraph, index):
    """Turn the index-th '[ ]' in the paragraph into '[X]'."""
    cnt = 0
    for t in _text_nodes(paragraph):
        s = t.text or ""
        if "[ ]" not in s:
            continue
        out, i = "", 0
        while True:
            j = s.find("[ ]", i)
            if j < 0:
                out += s[i:]
                break
            out += s[i:j] + ("[X]" if cnt == index else "[ ]")
            cnt += 1
            i = j + 3
        t.text = out


def _first(doc, needle, why):
    for p in doc.paragraphs:
        if needle in p.text:
            return p
    raise LeaseFillError(f"could not locate {why} (anchor {needle!r})")


def _all(doc, *needles):
    return [p for p in doc.paragraphs if all(n in p.text for n in needles)]


# ----------------------------------------------------------------------
# Template-aware filling
# ----------------------------------------------------------------------
_UTIL = {
    "city": ("Water Supply:", 0), "well": ("Water Supply:", 1),
    "sewer": ("Wastewater Disposal:", 0), "septic": ("Wastewater Disposal:", 1),
    "oil": ("Fuel:", 0), "gas": ("Fuel:", 1), "propane": ("Fuel:", 2),
}


def _fill_document(data):
    lease_type = data["lease_type"]
    doc = docx.Document(str(CONFIG["templates"][lease_type]))

    # Shared fields
    fill_blanks(_first(doc, "entered into on", "agreement date"), [data["agreement_date"]])
    landlord_p = next((p for p in doc.paragraphs if p.text.strip().startswith("Landlord:")), None)
    if landlord_p is None:
        raise LeaseFillError("could not locate landlord line")
    fill_blanks(landlord_p, [data["landlord"]])
    fill_blanks(_first(doc, "located at:", "premises address"), [data["premises_address"]])
    fill_blanks(_first(doc, "shall commence on", "lease term"), [data["term_start"], data["term_end"]])
    fill_blanks(_first(doc, "the sum of $", "monthly rent"), [data["rent"]])
    fill_blanks(_first(doc, "security deposit of $", "security deposit"), [data["deposit"]])

    # Tenants (two slots; fill one or two)
    blocks = _all(doc, "Address:", "City, State, Zip:")
    ssns = [p for p in doc.paragraphs if "Social Security Number" in p.text]
    if len(blocks) < 2 or len(ssns) < 2:
        raise LeaseFillError("expected two tenant slots in the template")
    for ti, t in enumerate(data["tenants"][:2]):
        if lease_type == "multi_family":
            fill_blanks(blocks[ti], [t["name"], t.get("address", ""), t.get("city_state_zip", "")])
        else:
            insert_before_break(blocks[ti], t["name"])
            fill_blanks(blocks[ti], [t.get("address", ""), t.get("city_state_zip", "")])
        if t.get("ssn"):
            fill_blanks(ssns[ti], [t["ssn"]])

    # Single-family utility checkboxes
    if lease_type == "single_family":
        for kind in ("water", "wastewater", "fuel"):
            val = (data.get("utilities", {}).get(kind) or "").strip().lower()
            if not val:
                continue
            if val not in _UTIL:
                raise LeaseFillError(f"unknown {kind} option {val!r} (expected {sorted(k for k,(a,_) in _UTIL.items() if a==_UTIL[val][0])})")
            anchor, idx = _UTIL[val]
            check_nth_box(_first(doc, anchor, f"{kind} checkbox"), idx)

    # Multifamily optional pets
    if lease_type == "multi_family" and data.get("pets"):
        pet_lines = [p for p in doc.paragraphs if "(Name / Breed / Color)" in p.text]
        for pi, pet in enumerate(data["pets"][:len(pet_lines)]):
            fill_blanks(pet_lines[pi], [pet])

    return doc


def _docx_to_pdf(docx_path, out_dir):
    """Convert with LibreOffice. soffice exits 0 even on failure, so success
    is decided by whether the PDF actually appeared."""
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    # A private profile forces an independent soffice instance, so conversion
    # works even when the LibreOffice Quickstarter is running (otherwise the
    # --convert-to is forwarded to the running instance and silently ignored).
    # .as_uri() yields a correct file:///C:/... URI on Windows (a bare
    # file://C:\path with backslashes will not parse).
    profile = (out_dir / f".lo-{docx_path.stem}").resolve()
    env = dict(os.environ)
    env.setdefault("HOME", str(out_dir))  # soffice needs a writable HOME on *nix
    cmd = [CONFIG["soffice_bin"], "--headless",
           f"-env:UserInstallation={profile.as_uri()}",
           "--convert-to", "pdf", "--outdir", str(out_dir), str(docx_path)]
    try:
        subprocess.run(cmd, env=env, capture_output=True,
                       timeout=CONFIG["convert_timeout_s"])
    except FileNotFoundError:
        raise LeaseFillError(f"LibreOffice not found ({CONFIG['soffice_bin']!r}). "
                             f"Install LibreOffice on this PC or set LEASE_SOFFICE.")
    except subprocess.TimeoutExpired:
        raise LeaseFillError("LibreOffice conversion timed out")
    pdf = out_dir / (docx_path.stem + ".pdf")
    if not pdf.exists() or pdf.stat().st_size == 0:
        raise LeaseFillError("DOCX->PDF conversion produced no PDF "
                             "(is LibreOffice working on this PC?)")
    return pdf


def fill_lease_pdf(data, out_pdf):
    """Fill the template and render it to out_pdf. Returns out_pdf."""
    out_pdf = Path(out_pdf)
    doc = _fill_document(data)
    tmp_docx = out_pdf.with_suffix(".docx")
    doc.save(str(tmp_docx))
    try:
        produced = _docx_to_pdf(tmp_docx, out_pdf.parent)
        if produced.resolve() != out_pdf.resolve():
            produced.replace(out_pdf)
    finally:
        tmp_docx.unlink(missing_ok=True)
    return out_pdf


# ----------------------------------------------------------------------
# Intake -> normalized data
# ----------------------------------------------------------------------
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_LEASE_TYPES = {"single_family", "multi_family"}


def _fmt_date(v):
    if not v:
        return date.today().strftime("%B %-d, %Y")
    for fmt in ("%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(v, fmt).strftime("%B %-d, %Y")
        except ValueError:
            pass
    return str(v)


def _iso_date(v):
    """YYYY-MM-DD for the filing filename, or '' if the format is unknown."""
    for fmt in ("%Y-%m-%d", "%m/%d/%Y"):
        try:
            return datetime.strptime(str(v), fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return ""


def normalize_intake(raw):
    if not isinstance(raw, dict):
        raise LeaseFillError("intake is not a JSON object")
    lease_type = raw.get("lease_type")
    if lease_type not in _LEASE_TYPES:
        raise LeaseFillError(f"lease_type must be one of {sorted(_LEASE_TYPES)}, got {lease_type!r}")
    tenants = raw.get("tenants") or []
    if not (1 <= len(tenants) <= 2):
        raise LeaseFillError("intake must have 1 or 2 tenants")
    for i, t in enumerate(tenants, 1):
        if not t.get("name"):
            raise LeaseFillError(f"tenant {i} missing name")
        if not t.get("email") or not _EMAIL_RE.match(t["email"]):
            raise LeaseFillError(f"tenant {i} missing/invalid email: {t.get('email')!r}")
    required = ["landlord", "property", "premises_address", "term_start", "term_end", "rent", "deposit"]
    missing = [k for k in required if not raw.get(k)]
    if missing:
        raise LeaseFillError(f"intake missing fields: {missing}")
    return {
        "lease_type": lease_type,
        "landlord": raw["landlord"],
        "property": raw["property"],
        "premises_address": raw["premises_address"],
        "agreement_date": _fmt_date(raw.get("agreement_date")),
        "term_start": _fmt_date(raw["term_start"]),
        "term_end": _fmt_date(raw["term_end"]),
        "rent": raw["rent"],
        "deposit": raw["deposit"],
        "utilities": raw.get("utilities", {}),
        "pets": raw.get("pets", []),
        "tenants": tenants,
        # Filing metadata (used by lease_filer when the signed lease returns).
        # property_key/unit should match lease_folders.json; optional but
        # recommended for clean filing — otherwise slugify(property) is used.
        "property_key": raw.get("property_key", ""),
        "unit": raw.get("unit", ""),
        "filing_address": raw.get("filing_address", ""),
        "term_start_iso": _iso_date(raw["term_start"]),
    }


def _slug(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def job_id_for(data):
    surname = data["tenants"][0]["name"].split()[-1]
    return f"{_slug(data['property'])}-{_slug(surname)}"[:80]


def build_job(data, pdf_path):
    """The job contract lease_watcher / lease_sender consume. NO SSN here."""
    signers = [{"name": t["name"], "email": t["email"]} for t in data["tenants"]]
    surname = data["tenants"][0]["name"].split()[-1]
    return {
        "property": data["property"],
        "signing_name": f"Lease - {data['property']} - {surname}",
        "signers": signers,
        "tenant_name": signers[0]["name"],    # legacy single-signer mirror
        "tenant_email": signers[0]["email"],
        "pdf_path": str(pdf_path),
        # Filing metadata for lease_filer (the signed-lease return trip).
        "property_key": data.get("property_key") or _slug(data["property"]),
        "unit": data.get("unit", ""),
        "filing_address": data.get("filing_address") or data["property"].split(",")[0].strip(),
        "term_start_iso": data.get("term_start_iso", ""),
    }


# ----------------------------------------------------------------------
# APPROVALS-RAILS CONTRACT — the approval card, matching the REAL supervisor
# shape (ANSWERS_LEASEAGENT.md). Hard requirements: `id` (== filename, and it
# encodes the job as lease-<slug> so the decision id maps back), `created_at`
# (ISO, the sort key), `title` (push body), a flat human-first `body`, and the
# `actions` array. `pdf_url` is the tap-to-open link. NO SSN anywhere here.
# NOTE: the phone app needs a `kind:"lease"` card path added before these
# render with the right buttons — that's an app-thread task (see the handoff).
# ----------------------------------------------------------------------
def card_id_for(job_id):
    return f"lease-{job_id}"


def build_card(data, job_id, pdf_path, pdf_url=""):
    tenants = "; ".join(f"{t['name']} <{t['email']}>" for t in data["tenants"])
    kind_label = "Single-family" if data["lease_type"] == "single_family" else "Multifamily"
    body = "\n".join([
        f"Property: {data['property']}",
        f"Type: {kind_label} lease",
        f"Tenant(s): {tenants}",
        f"Rent: ${data['rent']}/mo    Deposit: ${data['deposit']}",
        f"Term: {data['term_start']} – {data['term_end']}",
    ])
    rel = f"{CONFIG['dropbox_rel_root'].rstrip('/')}/Pending/{Path(pdf_path).name}"
    card = {
        "id": card_id_for(job_id),          # filename must equal this
        "agent": "lease",
        "kind": "lease",
        "created_at": datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace("+00:00", "Z"),
        "title": f"Lease ready: {data['property']}",
        "subject": f"Lease ready: {data['property']}",
        "body": body,
        "actions": ["send", "reject"],      # no feedback until the fill agent honors it
        "pdf_dropbox_path": rel,            # ignored by the supervisor; kept for reference
        "fields": {                         # rides along for a future richer card
            "lease_type": data["lease_type"],
            "property": data["property"],
            "tenants": [{"name": t["name"], "email": t["email"]} for t in data["tenants"]],
            "rent": f"${data['rent']}/mo",
            "deposit": f"${data['deposit']}",
            "term": f"{data['term_start']} – {data['term_end']}",
        },
    }
    if pdf_url:
        card["pdf_url"] = pdf_url
    return card


def _atomic_write(path, text):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    tmp = path.with_suffix(path.suffix + ".tmp")
    tmp.write_text(text, encoding="utf-8")
    tmp.replace(path)


def process_intake(intake_path, dry_run=False):
    raw = json.loads(Path(intake_path).read_text(encoding="utf-8"))
    data = normalize_intake(raw)
    job_id = job_id_for(data)

    if dry_run:
        out = Path(intake_path).with_name(f"{job_id}.pdf")
        fill_lease_pdf(data, out)
        print(f"[DRY-RUN] filled PDF written to {out} (nothing queued)")
        return out

    pending = CONFIG["pending_dir"]
    pending.mkdir(parents=True, exist_ok=True)
    # Fill to a temp name then move, so a half-written PDF is never seen in Pending.
    tmp_pdf = pending / f".{job_id}.pdf.tmp"
    fill_lease_pdf(data, tmp_pdf)
    pdf_path = pending / f"{job_id}.pdf"
    tmp_pdf.replace(pdf_path)

    job = build_job(data, pdf_path)
    _atomic_write(pending / f"{job_id}.json", json.dumps(job, indent=2))

    pdf_url = make_pdf_url(pdf_path)   # tap-to-open link (empty if no Dropbox token)

    # Card LAST: only advertise the job once PDF + job file are in place. Writing
    # the card into approvals\pending IS the phone notification (supervisor sweep).
    card = build_card(data, job_id, pdf_path, pdf_url)
    _atomic_write(CONFIG["pending_cards_dir"] / f"{card['id']}.json", json.dumps(card, indent=2))

    print(f"queued {card['id']}: {pdf_path}  (+ job file + approval card"
          f"{'' if pdf_url else '; NO pdf_url — set LEASE_DROPBOX_TOKEN'})")
    return job_id


def watch_intake(poll_seconds=5, settle_seconds=3):
    """Watch the intake folder and fill each intake JSON as it lands. This is
    the deployment mechanism: for now Jay (or any Claude chat) drops an
    intake.json here; when Samantha /chat lands, its lease skill writes the
    same file to the same folder and nothing here changes."""
    intake_dir = CONFIG["intake_dir"]
    processed = intake_dir / "_processed"
    failed = intake_dir / "_failed"
    for d in (intake_dir, processed, failed):
        d.mkdir(parents=True, exist_ok=True)
    print(f"LeaseFill watching {intake_dir} for *.json (Ctrl-C to stop)")
    try:
        while True:
            for p in sorted(intake_dir.glob("*.json")):
                try:
                    if time.time() - p.stat().st_mtime < settle_seconds:
                        continue
                except FileNotFoundError:
                    continue
                try:
                    process_intake(p)
                    p.replace(processed / p.name)
                except (LeaseFillError, Exception) as e:
                    print(f"intake {p.name} FAILED: {e}", file=sys.stderr)
                    push("Lease fill failed", f"{p.name}: {e}")
                    try:
                        p.replace(failed / p.name)
                    except Exception:
                        pass
            time.sleep(poll_seconds)
    except KeyboardInterrupt:
        print("LeaseFill watcher stopped.")


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--intake", help="fill one intake JSON")
    ap.add_argument("--watch", action="store_true", help="watch the intake folder")
    ap.add_argument("--dry-run", action="store_true", help="with --intake: build PDF only, queue nothing")
    args = ap.parse_args()
    if args.watch:
        watch_intake()
    elif args.intake:
        try:
            process_intake(args.intake, dry_run=args.dry_run)
        except LeaseFillError as e:
            print(f"LEASE FILL ABORTED: {e}", file=sys.stderr)
            sys.exit(1)
    else:
        ap.error("provide --intake <file> or --watch")
